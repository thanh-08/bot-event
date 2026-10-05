"""Xuất danh sách sự kiện ra file Excel theo tuần hoặc tháng.

  python main.py export week [YYYY-MM-DD]   7 ngày kể từ ngày đó (mặc định: hôm nay)
  python main.py export month [YYYY-MM]     cả tháng đó; không ghi thì lấy 30 ngày tới
Trên Telegram: /excel (tuần), /excel thang (tháng).
"""
import calendar, json, os, re
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import db
from common import PRIORITY_HIGH, PRIORITY_MID, PRIORITY_NAMES, PROVINCES, now_vn, priority_info

FONT = "Arial"
SHEET_DATA = "Danh sách sự kiện"
SHEET_SUM = "Tổng hợp"
HEADER_ROW = 4
LAST_ROW = 1000          # vùng công thức ở sheet Tổng hợp (đủ cho thêm tay vào danh sách)

COLUMNS = [  # (tiêu đề, độ rộng)
    ("STT", 6), ("Từ ngày", 12), ("Đến ngày", 12), ("Tỉnh/Thành", 17), ("Tên sự kiện", 46),
    ("Địa điểm", 30), ("Giờ", 9), ("Quy mô dự kiến (người)", 14), ("Pháo hoa", 9),
    ("Đại nhạc hội", 10), ("Mức ưu tiên", 13), ("Lý do ưu tiên", 44), ("Số nguồn tin", 9),
    ("Nguồn tin (link bài gốc)", 60),
]
COL_PROV, COL_PRIO = "D", "K"

PRIO_STYLE = {   # nền, chữ
    "Cao": ("FFC7CE", "9C0006"),
    "Trung bình": ("FFEB9C", "7F6000"),
    "Thấp": ("C6EFCE", "006100"),
}


def period(kind, arg=None, today=None):
    """Trả về (ngày đầu, ngày cuối, nhãn) của kỳ báo cáo."""
    today = today or now_vn().date()
    arg = (arg or "").strip()
    if kind == "week":
        try:
            start = date.fromisoformat(arg) if arg else today
        except ValueError:
            raise ValueError("Ngày ghi dạng YYYY-MM-DD, ví dụ 2026-10-12")
        end = start + timedelta(days=6)
        return start, end, f"Tuần {start:%d/%m}–{end:%d/%m/%Y}"
    if kind == "month":
        m = re.fullmatch(r"(\d{4})-(\d{1,2})", arg)
        if m:
            y, mo = int(m[1]), int(m[2])
            if not 1 <= mo <= 12:
                raise ValueError("Tháng phải từ 1 đến 12")
            start, end = date(y, mo, 1), date(y, mo, calendar.monthrange(y, mo)[1])
            return start, end, f"Tháng {mo:02d}/{y}"
        if arg:
            raise ValueError("Tháng ghi dạng YYYY-MM, ví dụ 2026-11")
        end = today + timedelta(days=29)
        return today, end, f"30 ngày tới ({today:%d/%m}–{end:%d/%m/%Y})"
    raise ValueError("Chỉ hỗ trợ: week hoặc month")


def _font(**kw):
    return Font(name=FONT, size=kw.pop("size", 10), **kw)


def build(kind, arg=None, out_path=None, c=None, today=None):
    """Tạo file Excel, trả về (đường dẫn file, số sự kiện)."""
    today = today or now_vn().date()
    start, end, label = period(kind, arg, today)
    own = c is None
    c = c or db.conn()
    try:
        evs = db.upcoming(c, start.isoformat(), end.isoformat())
    finally:
        if own:
            c.close()

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_DATA

    ws["A1"] = f"DANH SÁCH SỰ KIỆN MIỀN NAM — {label}"
    ws["A1"].font = _font(size=14, bold=True)
    ws["A2"] = (f"Xuất lúc {now_vn():%H:%M %d/%m/%Y} · {len(evs)} sự kiện · "
                "Mức ưu tiên tính theo quy mô và địa điểm (xem sheet Tổng hợp)")
    ws["A2"].font = _font(italic=True, color="595959")

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for i, (title, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(HEADER_ROW, i, title)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[HEADER_ROW].height = 32

    for n, ev in enumerate(evs, 1):
        r = HEADER_ROW + n
        label_p, _, why = priority_info(ev)
        pname = PRIORITY_NAMES[label_p]
        urls = json.loads(ev["urls"]) if isinstance(ev.get("urls"), str) else (ev.get("urls") or [])
        s_date = date.fromisoformat(ev["start_date"])
        e_date = date.fromisoformat(ev["end_date"]) if ev.get("end_date") else s_date
        vals = [n, s_date, e_date, ev["province"], ev["name"], ev.get("venue") or "", ev.get("time_text") or "",
                ev.get("attendance") or "Chưa rõ", "Có" if ev.get("fireworks") else "", "Có" if ev.get("concert") else "",
                pname, "; ".join(why), len(urls), urls[0] if urls else ""]
        for i, v in enumerate(vals, 1):
            cell = ws.cell(r, i, v)
            cell.font = _font()
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=i in (5, 6, 12),
                                       horizontal="center" if i in (1, 2, 3, 7, 9, 10, 11, 13) else None)
        ws.cell(r, 2).number_format = ws.cell(r, 3).number_format = "dd/mm/yyyy"
        ws.cell(r, 8).number_format = "#,##0"
        bg, fg = PRIO_STYLE[pname]
        pc = ws.cell(r, 11)
        pc.fill = PatternFill("solid", fgColor=bg)
        pc.font = _font(bold=True, color=fg)
        if urls:
            link = ws.cell(r, 14)
            link.hyperlink = urls[0]
            link.font = _font(color="0563C1", underline="single")

    last = HEADER_ROW + max(len(evs), 1)
    if not evs:
        ws.cell(HEADER_ROW + 1, 1, "Không có sự kiện nào trong kỳ này.").font = _font(italic=True)
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(COLUMNS))}{last}"
    ws.freeze_panes = f"A{HEADER_ROW + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"

    _summary(wb, label, len(evs))

    if not out_path:
        os.makedirs("exports", exist_ok=True)
        out_path = os.path.join("exports", f"su_kien_{'tuan' if kind == 'week' else 'thang'}_{start:%Y-%m-%d}.xlsx")
    wb.calculation.fullCalcOnLoad = True      # Excel tự tính lại công thức khi mở
    wb.save(out_path)
    return out_path, len(evs)


def _summary(wb, label, n_events):
    ws = wb.create_sheet(SHEET_SUM)
    q = f"'{SHEET_DATA}'"
    rng = lambda col: f"{q}!${col}${HEADER_ROW + 1}:${col}${LAST_ROW}"

    ws["A1"] = f"TỔNG HỢP SỰ KIỆN THEO TỈNH VÀ MỨC ƯU TIÊN — {label}"
    ws["A1"].font = _font(size=14, bold=True)
    ws["A2"] = "Số liệu tự tính từ sheet Danh sách sự kiện (công thức), sửa danh sách thì bảng tự cập nhật."
    ws["A2"].font = _font(italic=True, color="595959")

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    heads = ["Tỉnh/Thành", "Cao", "Trung bình", "Thấp", "Tổng"]
    for i, h in enumerate(heads, 1):
        cell = ws.cell(4, i, h)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
        cell.border = border
    for k, prov in enumerate(PROVINCES, 5):
        ws.cell(k, 1, prov)
        for j, pname in enumerate(["Cao", "Trung bình", "Thấp"], 2):
            ws.cell(k, j, f'=COUNTIFS({rng(COL_PROV)},$A{k},{rng(COL_PRIO)},{get_column_letter(j)}$4)')
        ws.cell(k, 5, f"=SUM(B{k}:D{k})")
    tot = 5 + len(PROVINCES)
    ws.cell(tot, 1, "Tổng cộng")
    for j in range(2, 6):
        col = get_column_letter(j)
        ws.cell(tot, j, f"=SUM({col}5:{col}{tot - 1})")
    for r in range(5, tot + 1):
        for j in range(1, 6):
            cell = ws.cell(r, j)
            cell.border = border
            cell.font = _font(bold=(r == tot))
            if j > 1:
                cell.alignment = Alignment(horizontal="center")
    for j, pname in enumerate(["Cao", "Trung bình", "Thấp"], 2):
        bg, fg = PRIO_STYLE[pname]
        ws.cell(4, j).fill = PatternFill("solid", fgColor=bg)
        ws.cell(4, j).font = _font(bold=True, color=fg)

    ws.cell(tot + 2, 1, f"Kiểm tra: tổng số sự kiện trong danh sách = ").font = _font(color="595959")
    ws.cell(tot + 2, 4, f"=COUNTA({rng(COL_PROV)})").font = _font(color="595959")
    ws.cell(tot + 2, 5, f'=IF(D{tot + 2}=E{tot},"khớp","LỆCH (có tỉnh ngoài danh sách)")').font = _font(color="595959")

    note = tot + 4
    ws.cell(note, 1, "Cách xếp mức ưu tiên đối với mạng lưới").font = _font(bold=True)
    rules = [
        "Cao: quy mô từ 10.000 người, có bắn pháo hoa, đại nhạc hội/đêm nhạc lớn, hoặc tổng điểm từ "
        f"{PRIORITY_HIGH} trở lên.",
        f"Trung bình: tổng điểm từ {PRIORITY_MID} đến {PRIORITY_HIGH - 1}.",
        "Thấp: các sự kiện còn lại.",
        "Điểm cộng theo quy mô: 1.000+ người (1), 5.000+ (2), 10.000+ (3), 20.000+ (4); pháo hoa (+2); "
        "đại nhạc hội (+2); kéo dài từ 3 ngày (+1); lễ hội truyền thống lớn như Vía Bà, Ok Om Bok, Nghinh Ông… (+3).",
        "Điểm cộng theo địa điểm: sân bay/cảng, khu du lịch đông khách, sân vận động, quảng trường… (1–3 điểm, "
        "lấy địa điểm cao nhất); TP. Hồ Chí Minh (+1).",
        "Đây là quy tắc tự đặt, chỉnh danh sách địa điểm và ngưỡng điểm trong file src/common.py "
        "(HOTSPOTS, BIG_FESTIVALS, PRIORITY_HIGH, PRIORITY_MID).",
    ]
    for i, t in enumerate(rules, 1):
        cell = ws.cell(note + i, 1, t)
        cell.font = _font()
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=note + i, start_column=1, end_row=note + i, end_column=5)
        ws.row_dimensions[note + i].height = 14 * (len(t) // 75 + 1) + 2     # bảng rộng khoảng 80 ký tự
    ws.column_dimensions["A"].width = 24
    for col in "BCDE":
        ws.column_dimensions[col].width = 14
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
